import io,json,re,secrets,zipfile
from uuid import uuid4
from urllib.parse import urlencode
import httpx
from fastapi import FastAPI,Request,UploadFile,File,HTTPException
from fastapi.middleware.cors import CORSMiddleware
from starlette.middleware.sessions import SessionMiddleware
from fastapi.responses import RedirectResponse,StreamingResponse
from pypdf import PdfReader
from docx import Document
from .config import *
from .schemas import Spec,Generate,RepoCreate,Publish
from .security import safe_files,safe_path,include
app=FastAPI(title="AI Project Builder")
app.add_middleware(SessionMiddleware,secret_key=SESSION_SECRET,same_site="lax",https_only=False,max_age=3600)
app.add_middleware(CORSMiddleware,allow_origins=[FRONTEND_ORIGIN],allow_credentials=True,allow_methods=["GET","POST"],allow_headers=["Content-Type"])
PROJECTS={}
def owner(r):
 if "owner" not in r.session:r.session["owner"]=secrets.token_urlsafe(24)
 return r.session["owner"]
def getp(r,pid):
 p=PROJECTS.get(pid)
 if not p or p["owner"]!=owner(r):raise HTTPException(404,"Project not found")
 return p
def gh(r):
 t=r.session.get("gh_token")
 if not t:raise HTTPException(401,"Connect GitHub first")
 return t
@app.get("/api/health")
def health():return {"status":"ok","ai_configured":bool(AI_API_KEY),"github_configured":bool(GITHUB_CLIENT_ID and GITHUB_CLIENT_SECRET)}
@app.post("/api/projects/analyze")
def analyze(r:Request,s:Spec):return {"analysis":{"stack":s.stack,"requirements":[x.strip(" -•") for x in s.topic.splitlines() if x.strip()][:30],"privacy_note":"Federated learning does not automatically guarantee privacy."},"ready":True}
@app.post("/api/projects/upload-spec")
async def upload(r:Request,file:UploadFile=File(...)):
 name=(file.filename or "file").replace("\\\\","/").split("/")[-1]; ext="."+name.rsplit(".",1)[-1].lower() if "." in name else ""
 if ext not in {".pdf",".docx",".txt",".png",".jpg",".jpeg",".webp"}:raise HTTPException(415,"Use PDF, DOCX, TXT, PNG, JPG or WEBP")
 data=await file.read(MAX_UPLOAD_MB*1024*1024+1)
 if len(data)>MAX_UPLOAD_MB*1024*1024:raise HTTPException(413,"Upload too large")
 if ext==".txt":text=data.decode("utf-8",errors="replace")
 elif ext==".pdf":
  try:text="\n".join(x.extract_text() or "" for x in PdfReader(io.BytesIO(data)).pages)
  except Exception:raise HTTPException(400,"Could not read PDF")
 elif ext==".docx":
  try:text="\n".join(x.text for x in Document(io.BytesIO(data)).paragraphs)
  except Exception:raise HTTPException(400,"Could not read DOCX")
 else:raise HTTPException(422,"Image OCR is not configured")
 if not text.strip():raise HTTPException(422,"No text extracted")
 return {"filename":name,"text":text[:20000]}
async def ai(topic,stack,name):
 if not AI_API_KEY:raise HTTPException(503,"Set AI_API_KEY in backend/.env")
 system='You are a senior engineer. Return only JSON {"files":{"relative/path":"complete text"},"summary":"short summary"}. Create a coherent runnable project with at most 60 text files and 1.5MB. Include README.md, .gitignore, .env.example placeholders, dependency manifests and tests. Never include .env, credentials, tokens, binaries, node_modules or build output. Medical AI is research-only, uses de-identified licensed data, and makes no clinical claims.'
 async with httpx.AsyncClient(timeout=120) as c:
  try:
   r=await c.post(f"{AI_BASE_URL}/chat/completions",headers={"Authorization":f"Bearer {AI_API_KEY}"},json={"model":AI_MODEL,"temperature":0.2,"response_format":{"type":"json_object"},"messages":[{"role":"system","content":system},{"role":"user","content":f"Project {name}\nStack {stack}\nRequirements:\n{topic}"}]})
   r.raise_for_status();d=json.loads(r.json()["choices"][0]["message"]["content"])
  except httpx.HTTPStatusError as e:raise HTTPException(502,f"AI provider HTTP {e.response.status_code}")
  except Exception as e:raise HTTPException(502,f"AI generation failed: {type(e).__name__}")
 fs=safe_files(d.get("files"));fs.setdefault("README.md",f"# {name}\nReview generated code before running.");fs.setdefault(".gitignore",".env\n.env.*\n!.env.example\nnode_modules/\ndist/\nbuild/\n");fs.setdefault(".env.example","# Add environment variables here; never commit real secrets.")
 return fs,str(d.get("summary","Generated source files"))[:1000]
@app.post("/api/projects/generate")
async def generate(r:Request,s:Generate):
 fs,summary=await ai(s.topic,s.stack,s.project_name);pid=str(uuid4());PROJECTS[pid]={"owner":owner(r),"name":s.project_name,"files":fs,"summary":summary}
 return {"project_id":pid,"name":s.project_name,"summary":summary,"files":[{"path":k,"size":len(v.encode())} for k,v in sorted(fs.items())]}
@app.get("/api/projects/{pid}/files")
def getfiles(r:Request,pid:str):
 p=getp(r,pid);return {"project_id":pid,"name":p["name"],"summary":p["summary"],"files":[{"path":k,"content":v,"size":len(v.encode())} for k,v in sorted(p["files"].items())]}
@app.get("/api/projects/{pid}/download")
def download(r:Request,pid:str):
 p=getp(r,pid);b=io.BytesIO()
 with zipfile.ZipFile(b,"w",zipfile.ZIP_DEFLATED) as z:
  for path,content in p["files"].items():
   path=safe_path(path)
   if include(path):z.writestr(f'{p["name"]}/{path}',content)
 b.seek(0);name=re.sub(r"[^A-Za-z0-9._-]","-",p["name"])[:80]
 return StreamingResponse(b,media_type="application/zip",headers={"Content-Disposition":f'attachment; filename="{name}.zip"'})
@app.get("/api/github/login")
def login(r:Request):
 if not GITHUB_CLIENT_ID or not GITHUB_CLIENT_SECRET:raise HTTPException(503,"Configure GitHub OAuth in backend/.env")
 state=secrets.token_urlsafe(32);r.session["oauth_state"]=state
 return {"url":"https://github.com/login/oauth/authorize?"+urlencode({"client_id":GITHUB_CLIENT_ID,"redirect_uri":GITHUB_CALLBACK_URL,"scope":"repo","state":state})}
@app.get("/api/github/callback")
async def callback(r:Request,code:str="",state:str="",error:str=""):
 if error:return RedirectResponse(FRONTEND_ORIGIN+"/?github_error=1")
 old=r.session.pop("oauth_state",None)
 if not old or not secrets.compare_digest(old,state):raise HTTPException(400,"Invalid OAuth state")
 async with httpx.AsyncClient(timeout=30) as c:x=await c.post("https://github.com/login/oauth/access_token",headers={"Accept":"application/json"},data={"client_id":GITHUB_CLIENT_ID,"client_secret":GITHUB_CLIENT_SECRET,"code":code,"redirect_uri":GITHUB_CALLBACK_URL})
 if x.status_code>=400 or not x.json().get("access_token"):raise HTTPException(502,"GitHub token exchange failed")
 r.session["gh_token"]=x.json()["access_token"]
 return RedirectResponse(FRONTEND_ORIGIN+"/?github_connected=1")
@app.get("/api/github/repositories")
async def repos(r:Request):
 h={"Authorization":f"Bearer {gh(r)}","Accept":"application/vnd.github+json"}
 async with httpx.AsyncClient(timeout=30) as c:x=await c.get("https://api.github.com/user/repos?per_page=100&sort=updated",headers=h)
 if x.status_code>=400:raise HTTPException(502,"Could not list repositories")
 return {"repositories":[{"full_name":a["full_name"],"private":a["private"],"default_branch":a["default_branch"],"html_url":a["html_url"]} for a in x.json()]}
@app.post("/api/github/repositories")
async def create_repo(r:Request,s:RepoCreate):
 h={"Authorization":f"Bearer {gh(r)}","Accept":"application/vnd.github+json"}
 async with httpx.AsyncClient(timeout=30) as c:x=await c.post("https://api.github.com/user/repos",headers=h,json={"name":s.name,"description":s.description,"private":s.private,"auto_init":True})
 if x.status_code==422:raise HTTPException(409,"Repository name conflict")
 if x.status_code>=400:raise HTTPException(502,"GitHub repository creation failed")
 a=x.json();return {"full_name":a["full_name"],"html_url":a["html_url"],"default_branch":a["default_branch"],"private":a["private"]}
@app.post("/api/github/publish-confirmation")
def confirm(r:Request):
 gh(r);t=secrets.token_urlsafe(32);r.session["publish_confirm"]=t;return {"confirmation_token":t}
@app.post("/api/github/publish")
async def publish(r:Request,s:Publish):
 token=gh(r);p=getp(r,s.project_id);old=r.session.pop("publish_confirm",None)
 if not old or not secrets.compare_digest(old,s.confirmation_token):raise HTTPException(400,"Confirmation expired")
 h={"Authorization":f"Bearer {token}","Accept":"application/vnd.github+json","X-GitHub-Api-Version":"2022-11-28"}
 async with httpx.AsyncClient(timeout=60) as c:
  rr=await c.get(f"https://api.github.com/repos/{s.repo_full_name}",headers=h)
  if rr.status_code>=400:raise HTTPException(404,"Repository unavailable")
  repo=rr.json();ref=await c.get(f"https://api.github.com/repos/{s.repo_full_name}/git/ref/heads/{s.branch}",headers=h)
  if ref.status_code==404:
   if repo.get("default_branch") and s.branch!=repo["default_branch"]:raise HTTPException(409,"Branch not found")
   parent=None
  elif ref.status_code>=400:raise HTTPException(502,"Could not read branch")
  else:parent=ref.json()["object"]["sha"]
  tree=[]
  for path,content in p["files"].items():
   path=safe_path(path)
   if not include(path):continue
   blob=await c.post(f"https://api.github.com/repos/{s.repo_full_name}/git/blobs",headers=h,json={"content":content,"encoding":"utf-8"})
   if blob.status_code>=400:raise HTTPException(502,f"Upload failed for {path}")
   tree.append({"path":path,"mode":"100644","type":"blob","sha":blob.json()["sha"]})
  tb={"tree":tree}
  if parent:tb["base_tree"]=parent
  tr=await c.post(f"https://api.github.com/repos/{s.repo_full_name}/git/trees",headers=h,json=tb)
  if tr.status_code>=400:raise HTTPException(502,"Tree creation failed")
  cb={"message":s.message,"tree":tr.json()["sha"]}
  if parent:cb["parents"]=[parent]
  cr=await c.post(f"https://api.github.com/repos/{s.repo_full_name}/git/commits",headers=h,json=cb)
  if cr.status_code>=400:raise HTTPException(502,"Commit creation failed")
  sha=cr.json()["sha"]
  if parent:u=await c.patch(f"https://api.github.com/repos/{s.repo_full_name}/git/refs/heads/{s.branch}",headers=h,json={"sha":sha,"force":False})
  else:u=await c.post(f"https://api.github.com/repos/{s.repo_full_name}/git/refs",headers=h,json={"ref":f"refs/heads/{s.branch}","sha":sha})
  if u.status_code>=400:raise HTTPException(409,"Commit created but branch update failed")
 return {"success":True,"repository_url":repo["html_url"],"commit_url":f'{repo["html_url"]}/commit/{sha}',"commit_sha":sha,"file_count":len(tree)}
